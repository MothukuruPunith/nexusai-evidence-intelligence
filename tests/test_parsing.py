import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas

from nexusai.core.config import Settings
from nexusai.documents.errors import DocumentError
from nexusai.documents.parser import MAX_BYTES, parse, validate_upload
from nexusai.documents.storage import parse_isolated
from nexusai.main import create_app


def pdf_bytes(text="Employees receive 20 annual leave days."):
    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawString(50, 780, text)
    canvas.save()
    return output.getvalue()


def docx_bytes():
    output = io.BytesIO()
    document = Document()
    document.add_paragraph("Employees receive 25 annual leave days.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Europe"
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "data,kind,expected,locator",
    [
        (b"First paragraph.\n\nSecond paragraph.", "txt", "Second paragraph.", "paragraph:2"),
        (b"month,revenue\nSeptember,85000\n", "csv", "month: September", "csv_record:2"),
        (pdf_bytes(), "pdf", "20 annual leave days", "page:1"),
        (docx_bytes(), "docx", "Region | Europe", "table:2/row:1"),
    ],
)
def test_all_formats_preserve_source_locations(data, kind, expected, locator):
    blocks = parse(data, kind)
    assert any(expected in b["text"] and b["locator"] == locator for b in blocks)


@pytest.mark.parametrize(
    "data,kind,code",
    [
        (b"a,a\n1,2", "csv", "invalid_csv_headers"),
        (b"a,b\n1", "csv", "inconsistent_csv_columns"),
        (b"hello\x00", "txt", "binary_text_rejected"),
        (b"\xff", "txt", "utf8_required"),
        (b"%PDF-broken", "pdf", "malformed_document"),
        (b"PKbroken", "docx", "malformed_document"),
        (b"   ", "txt", "no_extractable_text"),
    ],
)
def test_invalid_documents_fail_explicitly(data, kind, code):
    with pytest.raises(DocumentError, match=code):
        parse(data, kind)


def test_scanned_or_blank_pdf_requires_review():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(DocumentError, match="ocr_or_blank_page_review_required"):
        parse(output.getvalue(), "pdf")


def test_encrypted_pdf_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("private")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(DocumentError, match="encrypted_pdf"):
        parse(output.getvalue(), "pdf")


def test_docx_expansion_limit():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"a" * (21 * 1024 * 1024))
    with pytest.raises(DocumentError, match="archive_expansion_limit"):
        parse(output.getvalue(), "docx")


def test_size_extension_and_mime_limits():
    for data, name, mime, status in [
        (b"a", "bad.exe", "", 415),
        (b"a", "good.txt", "application/pdf", 415),
        (b"a" * (MAX_BYTES + 1), "large.txt", "text/plain", 413),
    ]:
        with pytest.raises(DocumentError) as error:
            validate_upload(data, name, mime)
        assert error.value.status == status


def test_isolated_parser_success(tmp_path):
    source = tmp_path / "fixture.txt"
    source.write_text("An independently parsed document.")
    assert parse_isolated(source, "txt")[0]["text"] == source.read_text()


def test_parser_timeout_is_explicit(monkeypatch, tmp_path):
    import subprocess

    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("parser", 20)

    monkeypatch.setattr(subprocess, "run", timeout)
    with pytest.raises(DocumentError, match="parser_timeout"):
        parse_isolated(tmp_path / "unused", "txt")


def test_upload_body_limit_applies_before_database_and_multipart_parsing():
    app = create_app(Settings(_env_file=None, database_url=None))
    with TestClient(app) as client:
        response = client.post("/documents/upload", content=b"a" * (6 * 1024 * 1024 + 1))
        assert response.status_code == 413
        assert response.json()["error"] == "request_too_large"
        assert response.headers["x-request-id"]


def test_search_requires_database_and_valid_input():
    with TestClient(create_app(Settings(_env_file=None, database_url=None))) as client:
        assert client.post("/documents/search", json={"query": "leave days"}).status_code == 503
        assert client.post("/documents/search", json={"query": "x", "k": 100}).status_code == 422
